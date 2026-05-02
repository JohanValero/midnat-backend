import io
import markdown
import html2text
from docx import Document
from xhtml2pdf import pisa
from sqlalchemy.orm import Session
from app.models import NovelChapter, Fragment, Novel
from bs4 import BeautifulSoup

def get_chapter_content(db: Session, chapter_id: int) -> str:
    fragments = (
        db.query(Fragment)
        .filter(Fragment.chapter_id == chapter_id)
        .order_by(Fragment.order)
        .all()
    )
    return "\n\n".join(f.content for f in fragments)

def html_to_markdown(html_content: str) -> str:
    h = html2text.HTML2Text()
    h.ignore_links = False
    h.body_width = 0  # No wrap
    return h.handle(html_content)

def export_chapter(db: Session, chapter_id: int, format: str) -> tuple[io.BytesIO, str, str]:
    chapter = db.query(NovelChapter).filter(NovelChapter.id == chapter_id).first()
    if not chapter:
        raise ValueError("Chapter not found")
        
    content_html = get_chapter_content(db, chapter_id)
    title = f"Capítulo {chapter.chapter_number}: {chapter.title}"
    filename_base = f"Capitulo_{chapter.chapter_number}_{chapter.title.replace(' ', '_')}"
    
    return _generate_file(title, content_html, format, filename_base)

def export_novel(db: Session, novel_id: int, format: str) -> tuple[io.BytesIO, str, str]:
    novel = db.query(Novel).filter(Novel.id == novel_id).first()
    if not novel:
        raise ValueError("Novel not found")
        
    chapters = (
        db.query(NovelChapter)
        .filter(NovelChapter.novel_id == novel_id)
        .order_by(NovelChapter.chapter_number)
        .all()
    )
    
    full_html = ""
    for chapter in chapters:
        chapter_content = get_chapter_content(db, chapter.id)
        full_html += f"<h1>Capítulo {chapter.chapter_number}: {chapter.title}</h1>"
        full_html += chapter_content
        
    title = novel.title
    filename_base = novel.title.replace(' ', '_')
    
    return _generate_file(title, full_html, format, filename_base)

def _generate_file(title: str, content_html: str, format: str, filename_base: str) -> tuple[io.BytesIO, str, str]:
    format = format.lower()
    
    if format == "markdown" or format == "md":
        md_content = f"# {title}\n\n"
        md_content += html_to_markdown(content_html)
        return io.BytesIO(md_content.encode("utf-8")), f"{filename_base}.md", "text/markdown"
        
    elif format == "docx":
        doc = Document()
        doc.add_heading(title, 0)
        
        soup = BeautifulSoup(content_html, "html.parser")
        for p in soup.find_all(['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
            if p.name.startswith('h'):
                level = int(p.name[1])
                doc.add_heading(p.get_text(), level=min(level, 9))
            else:
                doc.add_paragraph(p.get_text())
                
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream, f"{filename_base}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        
    elif format == "pdf":
        html_template = f"""
        <html>
        <head>
            <meta charset="UTF-8">
            <style>
                @page {{
                    margin: 2cm;
                }}
                body {{
                    font-family: Helvetica, Arial, sans-serif;
                    font-size: 11pt;
                    line-height: 1.6;
                    color: #333;
                }}
                h1 {{
                    text-align: center;
                    color: #000;
                    margin-bottom: 1cm;
                }}
                p {{
                    margin-bottom: 0.5cm;
                    text-align: justify;
                }}
            </style>
        </head>
        <body>
            <h1>{title}</h1>
            {content_html}
        </body>
        </html>
        """
        file_stream = io.BytesIO()
        pisa.CreatePDF(html_template, dest=file_stream)
        file_stream.seek(0)
        return file_stream, f"{filename_base}.pdf", "application/pdf"
        
    else:
        raise ValueError(f"Unsupported format: {format}")
